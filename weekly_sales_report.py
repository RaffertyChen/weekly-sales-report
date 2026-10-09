import os
import smtplib
from email.message import EmailMessage
from pathlib import Path
import sys

import anthropic
import pandas as pd
from dotenv import load_dotenv
from openpyxl import load_workbook
from openpyxl.styles import Font

folder = Path(__file__).parent
load_dotenv(folder / ".env")
client = anthropic.Anthropic()
GMAIL_USER = os.getenv("GMAIL_USER")
GMAIL_RECIPIENT = os.getenv("GMAIL_RECIPIENT")
GMAIL_PASSWORD = os.getenv("GMAIL_PASSWORD")

#alert system
def send_alert(problem):
    try:
        msg = EmailMessage()
        msg["From"] = GMAIL_USER
        msg["To"] = GMAIL_RECIPIENT
        msg["Subject"] = "Weekly Sales Report problem - ALERT"
        msg.set_content(f"Error: {problem}\n\nThe report was not sent.")
        with smtplib.SMTP_SSL("smtp.gmail.com", 465, timeout=10) as smtp:
            smtp.login(GMAIL_USER, GMAIL_PASSWORD)
            smtp.send_message(msg)
        print("send_alert: Email sent successfully")
    except Exception as e:
        print("send_alert: Failed to send email:", type(e).__name__, e)

#sorting data with pandas
try:
    orders = pd.read_excel(folder / "orders_export.xlsx", sheet_name=1, skiprows=3)
except Exception as e:
    problem = "could not find or read the orders_export.xlsx file, so no report was made"
    print(problem)
    send_alert(problem)
    sys.exit(1)
orders.columns = orders.columns.str.strip()

orders["Order ID"] = orders["Order ID"].str.strip()

orders["Order Date"] = pd.to_datetime(orders["Order Date"], format="mixed", errors="coerce")

orders["Unit Price (USD)"] = (orders["Unit Price (USD)"].astype(str).str.replace(r"[^\d.]", "", regex=True))
orders["Unit Price (USD)"] = pd.to_numeric(orders["Unit Price (USD)"], errors="coerce")

orders["Qty"] = orders["Qty"].astype(str).str.replace(r"[^\d]", "", regex=True)
orders["Qty"] = pd.to_numeric(orders["Qty"], errors="coerce")

orders["Sales"] = orders["Qty"] * orders["Unit Price (USD)"]

orders = orders[orders["Order ID"].str.startswith("HP-", na=False)]

orders["Customer"] = orders["Customer"].str.strip().str.replace("  ", " ").str.title()

orders["Status"] = orders["Status"].str.strip().str.replace(" ", "").str.title()

orders = orders[orders["Status"] != "Cancelled"]
orders = orders.drop_duplicates(subset="Order ID")

unreadable = orders[
    orders["Order Date"].isna()
    | orders["Unit Price (USD)"].isna()
    | orders["Qty"].isna()
]
for order_id in unreadable["Order ID"]:
    print("Warning: skipped", order_id, "because it can't be read")
orders = orders.drop(unreadable.index)
orders["Qty"] = orders["Qty"].astype(int)
orders = orders.drop(columns=["Status", "Line Total"])

total = orders["Sales"].sum()
by_month = orders.groupby(orders["Order Date"].dt.to_period("M"))["Sales"].sum()
by_customer = orders.groupby("Customer")["Sales"].sum().sort_values(ascending=False)

# AI API CODE
facts = f"""Orders: {len(orders)}
Total sales: ${total:,.2f}
Sales by month:
{by_month.to_string()}
Sales by customer:
{by_customer.to_string()}"""


instruction = """Write a 3-sentence summary of these sales numbers for a busy shop owner.
Only use numbers exactly as they appear below. Never add, subtract or combine numbers yourself.
Plain text only, no headings."""

try:
    message = client.messages.create(
        model="claude-haiku-4-5-20251001",
        max_tokens=300,
        messages=[{"role": "user", "content": instruction + "\n\n" + facts}],
    )
    summary = message.content[0].text
except Exception as e:
    print("Could not get the summary:", type(e).__name__, e)
    summary = "Summary not available"

by_month.index = by_month.index.astype(str)
by_month.index.name = "Month"

#workbook maker
report = folder / "sales_report.xlsx"
with pd.ExcelWriter(
    report, date_format="yyyy-mm-dd", datetime_format="yyyy-mm-dd"
) as writer:
    orders.to_excel(writer, sheet_name="Orders", index=False)
    by_customer.to_excel(writer, sheet_name="Sales by Customer", index=True)
    by_month.to_excel(writer, sheet_name="Sales by Month", index=True)


#workbook formatting
wb = load_workbook(report)
for ws in wb.worksheets:
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 24
for ws, letters in [(wb["Orders"], "FG"), (wb["Sales by Customer"], "B"), (wb["Sales by Month"], "B")]:
    for letter in letters:
        for cell in ws[letter][1:]:
            cell.number_format = "$#,##0.00"
wb.save(report)

#email
msg = EmailMessage()
msg["From"] = GMAIL_USER
msg["To"] = GMAIL_RECIPIENT
msg["Subject"] = "Weekly Sales Report"
msg.set_content(f"""Hi,

{summary}

The full report is attached.

Best regards,
Raff""")

with open(report, "rb") as f:
    msg.add_attachment(
        f.read(),
        maintype="application",
        subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=report.name,
    )
try:
    with smtplib.SMTP_SSL("smtp.gmail.com", 465, timeout=10) as smtp:
        smtp.login(GMAIL_USER, GMAIL_PASSWORD)
        smtp.send_message(msg)
    print("Email sent successfully")
except Exception as e:
    print("Failed to send email:", type(e).__name__, e)
    print("The report is still saved here:", report)
    sys.exit(1)
